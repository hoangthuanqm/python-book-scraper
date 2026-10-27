"""Book catalogue scraper for https://books.toscrape.com (a public sandbox site made for scraping practice).

Collects every book's title, price, rating, stock, category, UPC and URL,
then exports the result to CSV, JSON and a formatted Excel report.

Usage:
    python scraper.py                       # scrape everything
    python scraper.py --category travel     # one category
    python scraper.py --max-pages 3 --workers 8 --out output
"""

from __future__ import annotations

import argparse
import csv
import json
import logging
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import asdict, dataclass, fields
from pathlib import Path
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

BASE_URL = "https://books.toscrape.com/"
RATINGS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}

log = logging.getLogger("scraper")


@dataclass
class Book:
    title: str
    category: str
    price_gbp: float
    rating: int
    in_stock: int
    upc: str
    url: str


def make_session() -> requests.Session:
    """HTTP session with automatic retries and back-off for flaky networks."""
    session = requests.Session()
    session.headers["User-Agent"] = "Mozilla/5.0 (portfolio scraper; +https://books.toscrape.com)"
    retry = Retry(total=4, backoff_factor=0.5, status_forcelist=(429, 500, 502, 503, 504))
    session.mount("https://", HTTPAdapter(max_retries=retry, pool_maxsize=16))
    return session


def get_soup(session: requests.Session, url: str) -> BeautifulSoup:
    response = session.get(url, timeout=15)
    response.raise_for_status()
    response.encoding = "utf-8"
    return BeautifulSoup(response.text, "html.parser")


def find_category_url(session: requests.Session, name: str) -> str:
    soup = get_soup(session, BASE_URL)
    for link in soup.select("div.side_categories ul li ul li a"):
        if link.get_text(strip=True).lower() == name.lower():
            return urljoin(BASE_URL, link["href"])
    raise SystemExit(f"Category '{name}' not found.")


def collect_book_urls(session: requests.Session, start_url: str, max_pages: int | None) -> list[str]:
    """Walk listing pages via the 'next' link and return every book detail URL."""
    urls: list[str] = []
    page_url: str | None = start_url
    page = 0
    while page_url and (max_pages is None or page < max_pages):
        page += 1
        soup = get_soup(session, page_url)
        for a in soup.select("article.product_pod h3 a"):
            urls.append(urljoin(page_url, a["href"]))
        log.info("Listing page %d: %d books so far", page, len(urls))

        next_link = soup.select_one("li.next a")
        page_url = urljoin(page_url, next_link["href"]) if next_link else None
        time.sleep(0.2)  # be polite to the server
    return urls


def parse_book(session: requests.Session, url: str) -> Book:
    soup = get_soup(session, url)
    info = {row.th.get_text(strip=True): row.td.get_text(strip=True) for row in soup.select("table.table tr")}

    rating_class = soup.select_one("p.star-rating")["class"]
    rating = next((RATINGS[c] for c in rating_class if c in RATINGS), 0)

    stock_match = re.search(r"(\d+) available", info.get("Availability", ""))
    price = float(re.sub(r"[^\d.]", "", info.get("Price (incl. tax)", "0")))

    return Book(
        title=soup.select_one("div.product_main h1").get_text(strip=True),
        category=soup.select("ul.breadcrumb li a")[-1].get_text(strip=True),
        price_gbp=price,
        rating=rating,
        in_stock=int(stock_match.group(1)) if stock_match else 0,
        upc=info.get("UPC", ""),
        url=url,
    )


def scrape(category: str | None, max_pages: int | None, workers: int) -> list[Book]:
    session = make_session()
    start_url = find_category_url(session, category) if category else urljoin(BASE_URL, "catalogue/page-1.html")
    urls = collect_book_urls(session, start_url, max_pages)

    books: list[Book] = []
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = {pool.submit(parse_book, session, url): url for url in urls}
        for i, future in enumerate(as_completed(futures), 1):
            try:
                books.append(future.result())
            except Exception as exc:  # keep going if a single page fails
                log.warning("Failed %s: %s", futures[future], exc)
            if i % 50 == 0 or i == len(urls):
                log.info("Parsed %d/%d books", i, len(urls))

    books.sort(key=lambda b: (b.category, b.title))
    return books


def export_csv(books: list[Book], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=[f.name for f in fields(Book)])
        writer.writeheader()
        writer.writerows(asdict(b) for b in books)


def export_json(books: list[Book], path: Path) -> None:
    path.write_text(json.dumps([asdict(b) for b in books], indent=2, ensure_ascii=False), encoding="utf-8")


def export_excel(books: list[Book], path: Path) -> None:
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F46E5")

    ws = wb.active
    ws.title = "Books"
    headers = ["Title", "Category", "Price (£)", "Rating", "In stock", "UPC", "URL"]
    ws.append(headers)
    for b in books:
        ws.append([b.title, b.category, b.price_gbp, b.rating, b.in_stock, b.upc, b.url])

    for cell in ws[1]:
        cell.font, cell.fill = header_font, header_fill
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = "£#,##0.00"
    for col, width in enumerate([50, 20, 11, 8, 10, 18, 60], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Summary sheet: one row per category.
    summary = wb.create_sheet("Summary by category")
    summary.append(["Category", "Books", "Avg price (£)", "Avg rating", "Total stock"])
    by_category: dict[str, list[Book]] = {}
    for b in books:
        by_category.setdefault(b.category, []).append(b)
    for name, items in sorted(by_category.items(), key=lambda kv: -len(kv[1])):
        summary.append([
            name,
            len(items),
            round(sum(b.price_gbp for b in items) / len(items), 2),
            round(sum(b.rating for b in items) / len(items), 2),
            sum(b.in_stock for b in items),
        ])
    for cell in summary[1]:
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(horizontal="center")
    for col, width in enumerate([24, 10, 15, 12, 13], 1):
        summary.column_dimensions[get_column_letter(col)].width = width
    summary.freeze_panes = "A2"

    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Scrape books.toscrape.com and export to CSV / JSON / Excel.")
    parser.add_argument("--category", help="Only scrape this category, e.g. 'Travel'")
    parser.add_argument("--max-pages", type=int, help="Limit the number of listing pages (20 books each)")
    parser.add_argument("--workers", type=int, default=8, help="Parallel requests for detail pages (default 8)")
    parser.add_argument("--out", default="output", help="Output folder (default: output)")
    args = parser.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(asctime)s  %(levelname)-7s %(message)s", datefmt="%H:%M:%S")

    started = time.perf_counter()
    books = scrape(args.category, args.max_pages, args.workers)

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    export_csv(books, out / "books.csv")
    export_json(books, out / "books.json")
    export_excel(books, out / "books.xlsx")

    log.info("Done: %d books in %.1fs -> %s", len(books), time.perf_counter() - started, out.resolve())


if __name__ == "__main__":
    main()
